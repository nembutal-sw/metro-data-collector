from __future__ import annotations

import csv
import hashlib
import io
import re
import subprocess
import tempfile
import zipfile
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urljoin, urlparse

import numpy as np
from lxml import html  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]
from PIL import Image

from metro_collector.collectors.base import TransitSourceCollector
from metro_collector.collectors.http import ArtifactDownloader, ConditionalRequest
from metro_collector.collectors.kric_standard import _stable_station_code
from metro_collector.collectors.seoul_open_data import parse_service_seconds
from metro_collector.config import Settings
from metro_collector.domain.enums import (
    CoverageStatus,
    QualitySeverity,
    SourceAccessStatus,
    SourceKind,
)
from metro_collector.domain.models import (
    ActivationResult,
    CollectedArtifact,
    CollectionResult,
    NormalizedStopTime,
    QualityResult,
    SourceMetadata,
    ValidationResult,
)
from metro_collector.exceptions import SourceFormatError
from metro_collector.quality.validators import validate_normalized_stop_times

PASS_MARKERS = {"", "-", "--", "---"}
AREX_PASSENGER_STATIONS = (
    "서울",
    "공덕",
    "홍대입구",
    "디지털미디어시티",
    "마곡나루",
    "김포공항",
    "계양",
    "검암",
    "청라국제도시",
    "영종",
    "운서",
    "공항화물청사",
    "인천공항1터미널",
    "인천공항2터미널",
)
SERVICE_CODES = {
    "WEEKDAY": ("WEEKDAY",),
    "HOLIDAY": ("SATURDAY", "SUNDAY_HOLIDAY"),
}


@dataclass(frozen=True, slots=True)
class StationHtmlDataset:
    source_code: str
    source_name: str
    line_code: str
    effective_from: date
    station_order: tuple[str, ...]
    page_urls: tuple[str, ...]
    parser_kind: str
    forward_segment_seconds: tuple[int, ...]
    reverse_segment_seconds: tuple[int, ...]
    minimum_record_count: int = 500


@dataclass(slots=True)
class _SyntheticTrip:
    times: list[tuple[int, int]]


def _text(node: Any) -> str:
    return " ".join(node.text_content().split())


def _cell_seconds(value: object) -> int | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, (int, float)) and 0 <= float(value) < 1:
        return round(float(value) * 86400)
    text = str(value).strip()
    if text in PASS_MARKERS:
        return None
    return parse_service_seconds(text)


def _roll_forward(value: int | None, previous: int | None) -> int | None:
    if value is None:
        return None
    while previous is not None and value < previous - 12 * 3600:
        value += 24 * 3600
    return value


def iter_arex_workbook(path: Path) -> Iterator[NormalizedStopTime]:
    try:
        # The official workbook is a wide matrix. Normal mode avoids the very
        # expensive random-cell access behavior of openpyxl's read-only sheets.
        workbook = load_workbook(path, read_only=False, data_only=True)
    except Exception as error:
        raise SourceFormatError(f"Unable to open official AREX XLSX: {error}") from error
    try:
        sheet_map = {"평일": "WEEKDAY", "휴일": "HOLIDAY"}
        passenger_set = set(AREX_PASSENGER_STATIONS)
        for worksheet in workbook.worksheets:
            service_group = sheet_map.get(worksheet.title.strip())
            if service_group is None:
                continue
            blocks = (
                (4, 13, 41, "DOWN"),
                (46, 51, 79, "UP"),
            )
            for header_row, first_station_row, last_station_row, direction in blocks:
                for column in range(2, worksheet.max_column + 1):
                    train_number = str(worksheet.cell(header_row, column).value or "").strip()
                    if not train_number:
                        continue
                    raw_stops: list[tuple[str, int | None, int | None]] = []
                    previous: int | None = None
                    for station_row in range(first_station_row, last_station_row + 1, 2):
                        station_name = str(worksheet.cell(station_row, 1).value or "").strip()
                        if station_name not in passenger_set:
                            continue
                        arrival_value = worksheet.cell(station_row, column).value
                        departure_value = worksheet.cell(station_row + 1, column).value
                        arrival_text = str(arrival_value or "").strip()
                        departure_text = str(departure_value or "").strip()
                        if arrival_text in {"-", "--", "---"} or departure_text in {
                            "-",
                            "--",
                            "---",
                        }:
                            continue
                        arrival = _roll_forward(_cell_seconds(arrival_value), previous)
                        departure = _roll_forward(
                            _cell_seconds(departure_value),
                            arrival if arrival is not None else previous,
                        )
                        if arrival is None and departure is None:
                            continue
                        if arrival is None:
                            arrival = departure
                        if departure is None:
                            departure = arrival
                        previous = departure if departure is not None else arrival
                        raw_stops.append((station_name, arrival, departure))
                    if len(raw_stops) < 2:
                        continue
                    origin = raw_stops[0][0]
                    destination = raw_stops[-1][0]
                    express_type = "EXPRESS" if train_number.startswith("A1") else "LOCAL"
                    for service_code in SERVICE_CODES[service_group]:
                        trip_code = f"AREX:{service_code}:{train_number}"
                        for sequence, (station_name, arrival, departure) in enumerate(
                            raw_stops, start=1
                        ):
                            yield NormalizedStopTime(
                                line_code="AREX",
                                station_code=_stable_station_code("AREX", station_name),
                                station_name=station_name,
                                service_code=service_code,
                                direction=direction,
                                trip_code=trip_code,
                                train_number=train_number,
                                stop_sequence=sequence,
                                arrival_sec=arrival,
                                departure_sec=departure,
                                origin_name=origin,
                                destination_name=destination,
                                express_type=express_type,
                                source_row_number=header_row * 1000 + column,
                            )
    finally:
        workbook.close()


def _deterministic_zip(entries: dict[str, bytes], destination: Path) -> str:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = destination.with_suffix(destination.suffix + ".part")
    with zipfile.ZipFile(temporary_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, body in sorted(entries.items()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            archive.writestr(info, body)
    checksum = hashlib.sha256(temporary_path.read_bytes()).hexdigest()
    temporary_path.replace(destination)
    return checksum


def _latest_http_date(values: Sequence[str | None]) -> str | None:
    parsed: list[tuple[datetime, str]] = []
    for value in values:
        if not value:
            continue
        try:
            parsed.append((parsedate_to_datetime(value), value))
        except (TypeError, ValueError):
            continue
    if not parsed:
        return None
    return max(parsed)[1]


def _hourly_times(table: Any, *, hour_column: int, minute_column: int) -> list[int]:
    result: list[int] = []
    for row in table.xpath(".//tr"):
        cells = [_text(cell) for cell in row.xpath("./th|./td")]
        if len(cells) <= max(hour_column, minute_column):
            continue
        match = re.fullmatch(r"(\d{1,2})(?:시)?", cells[hour_column])
        if match is None:
            continue
        hour = int(match.group(1))
        for minute_text in re.findall(r"(?<!\d)(\d{2})(?!\d)", cells[minute_column]):
            minute = int(minute_text)
            if minute < 60:
                result.append(hour * 3600 + minute * 60)
    return sorted(set(result))


def _read_zip_pages(path: Path) -> list[bytes]:
    with zipfile.ZipFile(path) as archive:
        names = sorted(name for name in archive.namelist() if name.startswith("pages/"))
        return [archive.read(name) for name in names]


def _html_station_events(
    path: Path, dataset: StationHtmlDataset
) -> dict[tuple[str, str, int], list[int]]:
    pages = _read_zip_pages(path)
    if len(pages) != len(dataset.station_order):
        raise SourceFormatError(
            f"{dataset.source_code} expected {len(dataset.station_order)} station pages, "
            f"received {len(pages)}"
        )
    events: dict[tuple[str, str, int], list[int]] = {}
    for station_index, page in enumerate(pages):
        document = html.fromstring(page)
        tables = document.xpath("//table")
        if dataset.parser_kind == "UI_SINSEOL":
            if len(tables) < 4:
                raise SourceFormatError("Ui-Sinseol page no longer has four timetable tables")
            for service_group, table_index in (("WEEKDAY", 1), ("HOLIDAY", 3)):
                events[(service_group, "UP", station_index)] = _hourly_times(
                    tables[table_index], hour_column=1, minute_column=0
                )
                events[(service_group, "DOWN", station_index)] = _hourly_times(
                    tables[table_index], hour_column=1, minute_column=2
                )
        elif dataset.parser_kind == "SILLIM":
            if len(tables) < 6:
                raise SourceFormatError("Sillim page no longer has six timetable tables")
            for service_group, down_table, up_table in (
                ("WEEKDAY", 1, 2),
                ("HOLIDAY", 4, 5),
            ):
                events[(service_group, "DOWN", station_index)] = _hourly_times(
                    tables[down_table], hour_column=0, minute_column=1
                )
                events[(service_group, "UP", station_index)] = _hourly_times(
                    tables[up_table], hour_column=0, minute_column=1
                )
        else:
            raise SourceFormatError(f"Unknown station HTML parser: {dataset.parser_kind}")
    return events


def _segment_seconds(dataset: StationHtmlDataset, first: int, second: int) -> int:
    edge_index = min(first, second)
    values = (
        dataset.forward_segment_seconds
        if second > first
        else dataset.reverse_segment_seconds
    )
    return values[edge_index]


def _synthesize_trips(
    events: dict[tuple[str, str, int], list[int]],
    dataset: StationHtmlDataset,
    service_group: str,
    direction: str,
) -> list[_SyntheticTrip]:
    station_indices = (
        list(range(len(dataset.station_order)))
        if direction == "DOWN"
        else list(range(len(dataset.station_order) - 1, -1, -1))
    )
    all_trips: list[_SyntheticTrip] = []
    active: list[_SyntheticTrip] = []
    previous_index: int | None = None
    for position, station_index in enumerate(station_indices):
        station_events = events.get((service_group, direction, station_index), [])
        if position == len(station_indices) - 1 and not station_events and active:
            assert previous_index is not None
            expected = _segment_seconds(dataset, previous_index, station_index)
            for trip in active:
                trip.times.append((station_index, trip.times[-1][1] + expected))
            break
        if previous_index is None:
            active = [_SyntheticTrip(times=[(station_index, event)]) for event in station_events]
            all_trips.extend(active)
            previous_index = station_index
            continue
        expected = _segment_seconds(dataset, previous_index, station_index)
        available = set(range(len(active)))
        new_active: list[_SyntheticTrip] = []
        for event in station_events:
            candidates: list[tuple[int, int]] = []
            for candidate_index in available:
                last_time = active[candidate_index].times[-1][1]
                delta = event - last_time
                if -60 <= delta <= expected + 120:
                    candidates.append((abs(delta - expected), candidate_index))
            if candidates:
                _, candidate_index = min(candidates)
                available.remove(candidate_index)
                trip = active[candidate_index]
                adjusted = max(event, trip.times[-1][1] + min(30, expected))
                trip.times.append((station_index, adjusted))
            else:
                trip = _SyntheticTrip(times=[(station_index, event)])
                all_trips.append(trip)
            new_active.append(trip)
        active = new_active
        previous_index = station_index
    return [trip for trip in all_trips if len(trip.times) >= 2]


def iter_station_html_bundle(
    path: Path, dataset: StationHtmlDataset
) -> Iterator[NormalizedStopTime]:
    events = _html_station_events(path, dataset)
    for service_group in ("WEEKDAY", "HOLIDAY"):
        for direction in ("DOWN", "UP"):
            trips = _synthesize_trips(events, dataset, service_group, direction)
            for service_code in SERVICE_CODES[service_group]:
                for trip_index, trip in enumerate(trips, start=1):
                    origin = dataset.station_order[trip.times[0][0]]
                    destination = dataset.station_order[trip.times[-1][0]]
                    trip_code = (
                        f"{dataset.line_code}:{service_code}:{direction}:{trip_index:04d}:"
                        f"{trip.times[0][1]}"
                    )
                    for sequence, (station_index, event_time) in enumerate(
                        trip.times, start=1
                    ):
                        station_name = dataset.station_order[station_index]
                        yield NormalizedStopTime(
                            line_code=dataset.line_code,
                            station_code=_stable_station_code(dataset.line_code, station_name),
                            station_name=station_name,
                            service_code=service_code,
                            direction=direction,
                            trip_code=trip_code,
                            train_number=None,
                            stop_sequence=sequence,
                            arrival_sec=event_time,
                            departure_sec=event_time,
                            origin_name=origin,
                            destination_name=destination,
                            source_row_number=trip_index * 100 + sequence,
                        )


def _minimum_validation(
    records: Sequence[NormalizedStopTime], *, line_code: str, minimum: int
) -> ValidationResult:
    validation = validate_normalized_stop_times(records, required_line_codes={line_code})
    minimum_check = QualityResult(
        rule_code="EXPECTED_MINIMUM_RECORD_COUNT",
        severity=QualitySeverity.FATAL,
        passed=len(records) >= minimum,
        affected_count=max(0, minimum - len(records)),
        details={"parsed_rows": len(records), "minimum_rows": minimum},
    )
    return ValidationResult(
        passed=validation.passed and minimum_check.passed,
        checks=(*validation.checks, minimum_check),
    )


class ArexOfficialTimetableCollector(TransitSourceCollector):
    def __init__(self, settings: Settings) -> None:
        self._settings = settings
        self._downloader = ArtifactDownloader(
            timeout_seconds=settings.http_timeout_seconds,
            max_retries=settings.http_max_retries,
            timezone=settings.timezone,
        )
        self._cached_checksum: str | None = None
        self._cached_records: list[NormalizedStopTime] = []

    async def inspect(self) -> SourceMetadata:
        return SourceMetadata(
            code="AREX_TIMETABLE",
            name="공항철도 공식 전체 열차시간표",
            kind=SourceKind.DOWNLOAD_FILE,
            coverage=CoverageStatus.FULL_STATIC,
            access_status=SourceAccessStatus.ENABLED,
            base_url=self._settings.arex_official_timetable_page_url,
            stores_raw_artifact=True,
            notes="공항철도 공식 전체 XLSX를 매일 확인; 운영기관 TLS 호환 문제로 호스트 검증 예외",
        )

    async def _discover_url(self) -> str:
        override = self._settings.arex_official_timetable_download_url
        if override:
            return override
        page = await self._downloader.download(
            source_code="AREX_TIMETABLE_DISCOVERY",
            url=self._settings.arex_official_timetable_page_url,
            destination=self._settings.staging_data_dir / "arex-timetable-page.html",
            verify_tls=False,
        )
        if page is None:
            raise SourceFormatError("AREX timetable discovery page unexpectedly returned 304")
        document = html.fromstring(page.path.read_bytes())
        candidates: list[str] = []
        for raw_href in document.xpath("//a/@href|//button/@onclick"):
            href = str(raw_href)
            match = re.search(r"[\"']([^\"']+\.xlsx)[\"']", href, flags=re.IGNORECASE)
            file_href = match.group(1) if match else href
            if ".xlsx" in unquote(file_href).lower():
                candidates.append(
                    urljoin(self._settings.arex_official_timetable_page_url, file_href)
                )
        if not candidates:
            raise SourceFormatError("Official AREX page does not expose a timetable XLSX")

        def version_key(url: str) -> tuple[int, str]:
            digits = re.findall(r"(?<!\d)(\d{6,8})(?!\d)", unquote(url))
            return (max((int(value) for value in digits), default=0), url)

        return max(set(candidates), key=version_key)

    async def collect(self, previous: CollectedArtifact | None = None) -> CollectionResult | None:
        url = await self._discover_url()
        filename = Path(unquote(urlparse(url).path)).name or "arex-timetable.xlsx"
        use_conditional = previous is not None and previous.source_url == url
        artifact = await self._downloader.download(
            source_code="AREX_TIMETABLE",
            url=url,
            destination=self._settings.raw_data_dir / "AREX_TIMETABLE" / filename,
            conditional=ConditionalRequest(
                etag=previous.etag if use_conditional and previous else None,
                last_modified=previous.last_modified if use_conditional and previous else None,
            ),
            verify_tls=False,
        )
        if artifact is None or (previous and previous.checksum == artifact.checksum):
            return None
        return CollectionResult(
            artifact=artifact,
            effective_from=self._settings.arex_official_timetable_effective_from,
        )

    def _records(self, artifact: CollectedArtifact) -> list[NormalizedStopTime]:
        if self._cached_checksum != artifact.checksum:
            self._cached_records = list(iter_arex_workbook(artifact.path))
            self._cached_checksum = artifact.checksum
        return self._cached_records

    def iter_records(self, artifact: CollectedArtifact) -> Iterator[NormalizedStopTime]:
        yield from self._records(artifact)

    async def validate(self, artifact: CollectedArtifact) -> ValidationResult:
        return _minimum_validation(self._records(artifact), line_code="AREX", minimum=5000)

    async def activate(self, version_id: str) -> ActivationResult:
        return ActivationResult(
            activated=False,
            version_id=version_id,
            reason="Activation is transactionally performed by TimetableIngestionService",
        )


class HtmlStationTimetableCollector(TransitSourceCollector):
    def __init__(self, settings: Settings, dataset: StationHtmlDataset) -> None:
        self._settings = settings
        self._dataset = dataset
        self._downloader = ArtifactDownloader(
            timeout_seconds=settings.http_timeout_seconds,
            max_retries=settings.http_max_retries,
            timezone=settings.timezone,
        )
        self._cached_checksum: str | None = None
        self._cached_records: list[NormalizedStopTime] = []

    async def inspect(self) -> SourceMetadata:
        return SourceMetadata(
            code=self._dataset.source_code,
            name=self._dataset.source_name,
            kind=SourceKind.OFFICIAL_WEB,
            coverage=CoverageStatus.FULL_STATIC,
            access_status=SourceAccessStatus.ENABLED,
            base_url=self._dataset.page_urls[0],
            stores_raw_artifact=True,
            notes="운영기관 역별 HTML을 매일 묶음 수집하고 해시 변경 시에만 활성화",
        )

    async def collect(self, previous: CollectedArtifact | None = None) -> CollectionResult | None:
        entries: dict[str, bytes] = {}
        last_modified_values: list[str | None] = []
        for index, url in enumerate(self._dataset.page_urls, start=1):
            page = await self._downloader.download(
                source_code=f"{self._dataset.source_code}_PAGE_{index:02d}",
                url=url,
                destination=(
                    self._settings.staging_data_dir
                    / self._dataset.source_code
                    / f"station-{index:02d}.html"
                ),
                verify_tls=False,
            )
            if page is None:
                raise SourceFormatError("Station page unexpectedly returned 304 without page cache")
            entries[f"pages/station-{index:02d}.html"] = page.path.read_bytes()
            last_modified_values.append(page.last_modified)
        destination = (
            self._settings.raw_data_dir
            / self._dataset.source_code
            / "official-station-pages.zip"
        )
        checksum = _deterministic_zip(entries, destination)
        if previous and previous.checksum == checksum:
            return None
        artifact = CollectedArtifact(
            source_code=self._dataset.source_code,
            source_url=self._dataset.page_urls[0],
            retrieved_at=datetime.now(tz=self._settings.timezone),
            checksum=checksum,
            content_type="application/zip",
            path=destination,
            last_modified=_latest_http_date(last_modified_values),
        )
        return CollectionResult(artifact=artifact, effective_from=self._dataset.effective_from)

    def _records(self, artifact: CollectedArtifact) -> list[NormalizedStopTime]:
        if self._cached_checksum != artifact.checksum:
            self._cached_records = list(iter_station_html_bundle(artifact.path, self._dataset))
            self._cached_checksum = artifact.checksum
        return self._cached_records

    def iter_records(self, artifact: CollectedArtifact) -> Iterator[NormalizedStopTime]:
        yield from self._records(artifact)

    async def validate(self, artifact: CollectedArtifact) -> ValidationResult:
        return _minimum_validation(
            self._records(artifact),
            line_code=self._dataset.line_code,
            minimum=self._dataset.minimum_record_count,
        )

    async def activate(self, version_id: str) -> ActivationResult:
        return ActivationResult(
            activated=False,
            version_id=version_id,
            reason="Activation is transactionally performed by TimetableIngestionService",
        )


UIJEONGBU_STATIONS = (
    ("발곡", "balgok"),
    ("회룡", "hoeryong"),
    ("범골", "beomgol"),
    ("경전철의정부", "uijeongbu"),
    ("의정부시청", "city-hall"),
    ("흥선", "heungseon"),
    ("의정부중앙", "jungang"),
    ("동오", "dongo"),
    ("새말", "saemal"),
    ("경기도청북부청사", "north-office"),
    ("효자", "hyoja"),
    ("곤제", "gonje"),
    ("어룡(용현산업단지)", "eoryong"),
    ("송산", "songsan"),
    ("탑석", "tapseok"),
)


def _tesseract_minutes(image_body: bytes) -> tuple[list[int], list[int]]:
    with Image.open(io.BytesIO(image_body)) as source:
        image = source.convert("RGB")
    width, height = image.size
    data_top = round(height * 0.207)
    data_bottom = round(height * 0.814)
    pixels = np.asarray(image)
    region = pixels[data_top:data_bottom]
    green = (
        (region[:, :, 1] > 65)
        & (region[:, :, 1] > region[:, :, 0] * 1.35)
        & (region[:, :, 1] > region[:, :, 2] * 1.15)
    )
    columns = np.flatnonzero(green.sum(axis=0) > (data_bottom - data_top) * 0.45)
    groups: list[tuple[int, int]] = []
    for column in columns.tolist():
        if not groups or column > groups[-1][1] + 1:
            groups.append((column, column))
        else:
            groups[-1] = (groups[-1][0], column)
    groups = [group for group in groups if group[1] - group[0] >= 8]
    if len(groups) < 2:
        raise SourceFormatError("Unable to locate both timetable hour columns in ULRT image")
    first_hour, second_hour = groups[0], groups[1]
    crop = image.crop((0, data_top, width, data_bottom))
    crop_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as temporary:
            crop_path = Path(temporary.name)
        crop.save(crop_path)
        command = [
            "tesseract",
            str(crop_path),
            "stdout",
            "-l",
            "eng",
            "--psm",
            "6",
            "-c",
            "tessedit_char_whitelist=0123456789",
            "tsv",
        ]
        try:
            completed = subprocess.run(
                command,
                check=True,
                capture_output=True,
                text=True,
                encoding="utf-8",
                timeout=90,
            )
        except (FileNotFoundError, subprocess.SubprocessError) as error:
            raise SourceFormatError(f"ULRT timetable OCR failed: {error}") from error
    finally:
        if crop_path is not None:
            crop_path.unlink(missing_ok=True)
    schedules: tuple[dict[int, set[int]], dict[int, set[int]]] = ({}, {})
    reader = csv.DictReader(io.StringIO(completed.stdout), delimiter="\t")
    row_height = (data_bottom - data_top) / 20
    for record in reader:
        raw = re.sub(r"\D", "", record.get("text", ""))
        if not raw:
            continue
        try:
            left = int(record["left"])
            top = int(record["top"])
            box_width = int(record["width"])
            box_height = int(record["height"])
        except (KeyError, TypeError, ValueError):
            continue
        center_x = left + box_width / 2
        center_y = top + box_height / 2
        row_index = int(center_y / row_height)
        if not 0 <= row_index < 20:
            continue
        if first_hour[1] + 5 < center_x < second_hour[0] - 5:
            panel = 0
        elif second_hour[1] + 5 < center_x < width - 5:
            panel = 1
        else:
            continue
        tokens = [raw]
        if len(raw) > 2 and len(raw) % 2 == 0:
            tokens = [raw[index : index + 2] for index in range(0, len(raw), 2)]
        for token in tokens:
            minute = int(token.zfill(2))
            if 0 <= minute < 60:
                schedules[panel].setdefault(row_index, set()).add(minute)
    result: list[list[int]] = [[], []]
    for panel, rows in enumerate(schedules):
        for row_index, minutes in sorted(rows.items()):
            hour = 5 + row_index
            result[panel].extend(hour * 3600 + minute * 60 for minute in sorted(minutes))
    if len(result[0]) < 50 or len(result[1]) < 50:
        raise SourceFormatError(
            "ULRT timetable OCR returned too few departures: "
            f"weekday={len(result[0])}, holiday={len(result[1])}"
        )
    return result[0], result[1]


def _uijeongbu_dataset(effective_from: date) -> StationHtmlDataset:
    return StationHtmlDataset(
        source_code="UIJEONGBU_WEB",
        source_name="의정부경전철 공식 역별 시간표",
        line_code="UIJEONGBU_LRT",
        effective_from=effective_from,
        station_order=tuple(station for station, _ in UIJEONGBU_STATIONS),
        page_urls=(),
        parser_kind="UIJEONGBU_OCR",
        forward_segment_seconds=(60, 60, 60, 120, 60, 60, 120, 60, 120, 60, 120, 120, 60, 60),
        reverse_segment_seconds=(60, 60, 120, 60, 60, 120, 60, 60, 120, 60, 120, 120, 60, 120),
        minimum_record_count=5000,
    )


class UijeongbuOfficialTimetableCollector(TransitSourceCollector):
    def __init__(self, settings: Settings) -> None:
        self._settings = settings
        self._dataset = _uijeongbu_dataset(settings.uijeongbu_official_effective_from)
        self._downloader = ArtifactDownloader(
            timeout_seconds=settings.http_timeout_seconds,
            max_retries=settings.http_max_retries,
            timezone=settings.timezone,
        )
        self._cached_checksum: str | None = None
        self._cached_records: list[NormalizedStopTime] = []

    async def inspect(self) -> SourceMetadata:
        return SourceMetadata(
            code="UIJEONGBU_WEB",
            name="의정부경전철 공식 역별 시간표",
            kind=SourceKind.OFFICIAL_WEB,
            coverage=CoverageStatus.FULL_STATIC,
            access_status=SourceAccessStatus.ENABLED,
            base_url=self._settings.uijeongbu_official_timetable_page_url,
            stores_raw_artifact=True,
            notes="공식 역별 이미지 변경 감지 후 숫자 OCR; 품질 기준 통과 시에만 활성화",
        )

    async def collect(self, previous: CollectedArtifact | None = None) -> CollectionResult | None:
        page = await self._downloader.download(
            source_code="UIJEONGBU_WEB_PAGE",
            url=self._settings.uijeongbu_official_timetable_page_url,
            destination=self._settings.staging_data_dir / "UIJEONGBU_WEB" / "page.html",
        )
        if page is None:
            raise SourceFormatError("ULRT timetable page unexpectedly returned 304")
        document = html.fromstring(page.path.read_bytes())
        discovered = {
            Path(urlparse(urljoin(page.source_url, str(src))).path).name: urljoin(
                page.source_url, str(src)
            )
            for src in document.xpath('//img[contains(@src,"/timetable/")]/@src')
        }
        required = {
            f"{stem}-{suffix}.jpg"
            for _, stem in UIJEONGBU_STATIONS
            for suffix in ((1,) if stem == "balgok" else (1, 2))
        }
        missing = sorted(required - discovered.keys())
        if missing:
            raise SourceFormatError(f"ULRT page is missing timetable images: {missing}")
        entries = {"page.html": page.path.read_bytes()}
        last_modified_values: list[str | None] = [page.last_modified]
        for filename in sorted(required):
            image_artifact = await self._downloader.download(
                source_code=f"UIJEONGBU_IMAGE_{filename}",
                url=discovered[filename],
                destination=(
                    self._settings.staging_data_dir / "UIJEONGBU_WEB" / "images" / filename
                ),
            )
            if image_artifact is None:
                raise SourceFormatError("ULRT image unexpectedly returned 304 without image cache")
            entries[f"images/{filename}"] = image_artifact.path.read_bytes()
            last_modified_values.append(image_artifact.last_modified)
        destination = self._settings.raw_data_dir / "UIJEONGBU_WEB" / "official-images.zip"
        checksum = _deterministic_zip(entries, destination)
        if previous and previous.checksum == checksum:
            return None
        artifact = CollectedArtifact(
            source_code="UIJEONGBU_WEB",
            source_url=self._settings.uijeongbu_official_timetable_page_url,
            retrieved_at=datetime.now(tz=self._settings.timezone),
            checksum=checksum,
            content_type="application/zip",
            path=destination,
            last_modified=_latest_http_date(last_modified_values),
        )
        return CollectionResult(
            artifact=artifact,
            effective_from=self._settings.uijeongbu_official_effective_from,
        )

    def _records(self, artifact: CollectedArtifact) -> list[NormalizedStopTime]:
        if self._cached_checksum == artifact.checksum:
            return self._cached_records
        events: dict[tuple[str, str, int], list[int]] = {}
        with zipfile.ZipFile(artifact.path) as archive:
            for station_index, (_, stem) in enumerate(UIJEONGBU_STATIONS):
                if station_index < len(UIJEONGBU_STATIONS) - 1:
                    weekday, holiday = _tesseract_minutes(
                        archive.read(f"images/{stem}-1.jpg")
                    )
                    events[("WEEKDAY", "DOWN", station_index)] = weekday
                    events[("HOLIDAY", "DOWN", station_index)] = holiday
                if station_index > 0:
                    weekday, holiday = _tesseract_minutes(
                        archive.read(f"images/{stem}-2.jpg")
                    )
                    events[("WEEKDAY", "UP", station_index)] = weekday
                    events[("HOLIDAY", "UP", station_index)] = holiday
        records: list[NormalizedStopTime] = []
        for service_group in ("WEEKDAY", "HOLIDAY"):
            for direction in ("DOWN", "UP"):
                trips = _synthesize_trips(events, self._dataset, service_group, direction)
                for service_code in SERVICE_CODES[service_group]:
                    for trip_index, trip in enumerate(trips, start=1):
                        origin = self._dataset.station_order[trip.times[0][0]]
                        destination = self._dataset.station_order[trip.times[-1][0]]
                        trip_code = (
                            f"UIJEONGBU_LRT:{service_code}:{direction}:{trip_index:04d}:"
                            f"{trip.times[0][1]}"
                        )
                        for sequence, (station_index, event_time) in enumerate(
                            trip.times, start=1
                        ):
                            station_name = self._dataset.station_order[station_index]
                            records.append(
                                NormalizedStopTime(
                                    line_code="UIJEONGBU_LRT",
                                    station_code=_stable_station_code(
                                        "UIJEONGBU_LRT", station_name
                                    ),
                                    station_name=station_name,
                                    service_code=service_code,
                                    direction=direction,
                                    trip_code=trip_code,
                                    train_number=None,
                                    stop_sequence=sequence,
                                    arrival_sec=event_time,
                                    departure_sec=event_time,
                                    origin_name=origin,
                                    destination_name=destination,
                                    source_row_number=trip_index * 100 + sequence,
                                )
                            )
        self._cached_records = records
        self._cached_checksum = artifact.checksum
        return records

    def iter_records(self, artifact: CollectedArtifact) -> Iterator[NormalizedStopTime]:
        yield from self._records(artifact)

    async def validate(self, artifact: CollectedArtifact) -> ValidationResult:
        return _minimum_validation(
            self._records(artifact), line_code="UIJEONGBU_LRT", minimum=5000
        )

    async def activate(self, version_id: str) -> ActivationResult:
        return ActivationResult(
            activated=False,
            version_id=version_id,
            reason="Activation is transactionally performed by TimetableIngestionService",
        )


def build_official_operator_collectors(settings: Settings) -> list[TransitSourceCollector]:
    ui_stations = (
        "북한산우이",
        "솔밭공원",
        "4.19민주묘지",
        "가오리",
        "화계",
        "삼양",
        "삼양사거리",
        "솔샘",
        "북한산보국문",
        "정릉",
        "성신여대입구",
        "보문",
        "신설동",
    )
    sillim_stations = (
        "샛강",
        "대방",
        "서울지방병무청",
        "보라매",
        "보라매공원",
        "보라매병원",
        "당곡",
        "신림",
        "서원",
        "서울대벤처타운",
        "관악산",
    )
    ui_dataset = StationHtmlDataset(
        source_code="UI_SINSEOL_WEB",
        source_name="우이신설선 공식 역별 시간표",
        line_code="UI_SINSEOL",
        effective_from=settings.ui_sinseol_official_effective_from,
        station_order=ui_stations,
        page_urls=tuple(
            settings.ui_sinseol_official_page_template.format(station=index)
            for index in range(1, len(ui_stations) + 1)
        ),
        parser_kind="UI_SINSEOL",
        forward_segment_seconds=(60, 60, 60, 60, 60, 120, 60, 120, 120, 60, 60, 120),
        reverse_segment_seconds=(120, 60, 60, 120, 60, 60, 60, 120, 120, 60, 60, 60),
        minimum_record_count=5000,
    )
    sillim_dataset = StationHtmlDataset(
        source_code="SILLIM_WEB",
        source_name="신림선 공식 역별 시간표",
        line_code="SILLIM",
        effective_from=settings.sillim_official_effective_from,
        station_order=sillim_stations,
        page_urls=tuple(
            settings.sillim_official_page_template.format(station=index)
            for index in range(1, len(sillim_stations) + 1)
        ),
        parser_kind="SILLIM",
        forward_segment_seconds=(60, 60, 60, 120, 60, 60, 60, 60, 120, 120),
        reverse_segment_seconds=(60, 60, 60, 60, 60, 60, 60, 60, 120, 60),
        minimum_record_count=4000,
    )
    return [
        ArexOfficialTimetableCollector(settings),
        HtmlStationTimetableCollector(settings, ui_dataset),
        HtmlStationTimetableCollector(settings, sillim_dataset),
        UijeongbuOfficialTimetableCollector(settings),
    ]
