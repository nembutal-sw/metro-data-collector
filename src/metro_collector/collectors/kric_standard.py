from __future__ import annotations

import re
import statistics
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from itertools import pairwise
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from metro_collector.collectors.base import TransitSourceCollector
from metro_collector.collectors.http import ArtifactDownloader, ConditionalRequest
from metro_collector.collectors.seoul_open_data import parse_service_seconds
from metro_collector.config import Settings
from metro_collector.domain.enums import (
    CoverageStatus,
    SourceAccessStatus,
    SourceKind,
)
from metro_collector.domain.models import (
    ActivationResult,
    CollectedArtifact,
    CollectionResult,
    NormalizedStopTime,
    SourceMetadata,
    ValidationResult,
)
from metro_collector.exceptions import SourceFormatError
from metro_collector.quality.validators import validate_normalized_stop_times

KRIC_DATASET_PAGE = "https://openapi.kric.go.kr/rips/M_01_01/intro.do?lcd=A"
HEADER_ALIASES = {
    "train_number": ("열차번호",),
    "line_number": ("노선번호",),
    "line_name": ("노선명",),
    "origin": ("운행구간기점명",),
    "destination": ("운행구간종점명",),
    "operation_type": ("운행유형",),
    "service": ("요일구분",),
    "stations": ("운행구간정거장",),
    "arrivals": ("정거장도착시각", "정거장도착시간", "정가장도착시각"),
    # Some current KRIC workbooks contain the provider typo "정가장".
    "departures": ("정거장출발시각", "정거장출발시간", "정가장출발시각"),
    "data_as_of": ("데이터기준일자",),
}
HEADER_LOOKUP = {
    "".join(alias.split()): canonical
    for canonical, aliases in HEADER_ALIASES.items()
    for alias in aliases
}
REQUIRED_HEADERS = {"train_number", "service", "stations", "departures"}
SECONDS_PER_DAY = 24 * 60 * 60
ROLLOVER_THRESHOLD_SECONDS = 12 * 60 * 60


@dataclass(frozen=True, slots=True)
class KricDataset:
    source_code: str
    source_name: str
    dataset_id: int
    line_code: str
    effective_from: date
    download_url: str
    station_order: tuple[str, ...]
    minimum_record_count: int = 500
    fallback_segment_seconds: int = 90


@dataclass(slots=True)
class _ParsedStop:
    source_code: str
    station_name: str
    arrival_sec: int | None
    departure_sec: int | None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    return "".join(_cell_text(value).replace("\ufeff", "").split())


def normalize_station_name(value: str) -> str:
    normalized = " ".join(value.strip().split())
    if normalized.endswith("역") and len(normalized) > 1:
        normalized = normalized[:-1]
    return normalized


def _station_lookup_name(value: str) -> str:
    return re.sub(r"[\s·ㆍ._()\uFF08\uFF09-]", "", normalize_station_name(value)).upper()


def _stable_station_code(line_code: str, station_name: str) -> str:
    return f"{line_code}:{_station_lookup_name(station_name)}"


def parse_station_list(value: str) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    for token in value.split("+"):
        token = token.strip()
        if not token:
            continue
        try:
            code, name = token.split("-", 1)
        except ValueError as error:
            raise SourceFormatError(f"Invalid KRIC station item: {token!r}") from error
        station_name = normalize_station_name(name)
        if not code.strip() or not station_name:
            raise SourceFormatError(f"Invalid KRIC station item: {token!r}")
        result.append((code.strip(), station_name))
    if len(result) < 2:
        raise SourceFormatError("A KRIC trip must contain at least two stations")
    return result


def parse_timed_stops(value: str) -> dict[str, int | None]:
    text = value.strip()
    if not text:
        return {}
    tokens = [token.strip() for token in re.split(r"[+/]", text) if token.strip()]
    result: dict[str, int | None] = {}
    if tokens and all("-" in token for token in tokens):
        pairs = (token.split("-", 1) for token in tokens)
    else:
        if len(tokens) % 2:
            raise SourceFormatError(f"Invalid KRIC stop-time list: {value!r}")
        pairs = ((tokens[index], tokens[index + 1]) for index in range(0, len(tokens), 2))
    for raw_code, raw_time in pairs:
        code = raw_code.strip()
        time_text = raw_time.strip()
        if not code:
            raise SourceFormatError(f"Empty station code in KRIC stop-time list: {value!r}")
        if not time_text or time_text == ":":
            result[code] = None
        else:
            # A current Sillim workbook has one terminal time written as
            # ":0:59" between 0:57 and the end of service. Treat the leading
            # punctuation as a provider typo while preserving the intended 0:59.
            if re.fullmatch(r":\d{1,2}:\d{2}", time_text):
                time_text = time_text[1:]
            result[code] = parse_service_seconds(time_text)
    return result


def service_codes(value: str) -> tuple[str, ...]:
    normalized = "".join(value.upper().split())
    if normalized in {"평일", "주중", "WEEKDAY", "DAY"}:
        return ("WEEKDAY",)
    if normalized in {"토요일", "토", "SATURDAY", "SAT"}:
        return ("SATURDAY",)
    if normalized == "공휴일":
        return ("SUNDAY_HOLIDAY",)
    if normalized in {"휴일", "주말", "일요일", "HOLIDAY", "SUNDAY", "END"}:
        # KRIC operators generally publish one holiday table that is also used
        # on Saturdays. Store it against both application calendars.
        return ("SATURDAY", "SUNDAY_HOLIDAY")
    if "+" in normalized:
        parts = set(normalized.split("+"))
        has_saturday = bool(parts & {"토", "토요일", "주말", "SAT", "SATURDAY"})
        has_holiday = bool(
            parts & {"휴일", "일요일", "공휴일", "주말", "HOLIDAY", "SUNDAY"}
        )
        if has_saturday and has_holiday:
            return ("SATURDAY", "SUNDAY_HOLIDAY")
    raise SourceFormatError(f"Unknown KRIC service day: {value!r}")


def _direction(stations: Sequence[_ParsedStop], station_order: Sequence[str]) -> str:
    order = {_station_lookup_name(name): index for index, name in enumerate(station_order)}
    first = order.get(_station_lookup_name(stations[0].station_name))
    last = order.get(_station_lookup_name(stations[-1].station_name))
    if first is None or last is None or first == last:
        raise SourceFormatError(
            "KRIC trip endpoint is absent from the configured line order: "
            f"{stations[0].station_name!r} -> {stations[-1].station_name!r}"
        )
    return "DOWN" if first < last else "UP"


def _column_violations(stops: Sequence[_ParsedStop], attribute: str) -> tuple[int, int]:
    offset = 0
    previous: int | None = None
    violations = 0
    count = 0
    for stop in stops:
        value = getattr(stop, attribute)
        if value is None:
            continue
        count += 1
        if previous is not None and value + offset < previous - ROLLOVER_THRESHOLD_SECONDS:
            offset += SECONDS_PER_DAY
        adjusted = value + offset
        if previous is not None and adjusted < previous:
            violations += 1
        previous = adjusted
    return violations, count


def _prefer_consistent_time_column(stops: list[_ParsedStop]) -> None:
    arrival_violations, arrival_count = _column_violations(stops, "arrival_sec")
    departure_violations, departure_count = _column_violations(stops, "departure_sec")
    if departure_violations and not arrival_violations and arrival_count >= 2:
        for stop in stops:
            stop.departure_sec = None
    elif arrival_violations and not departure_violations and departure_count >= 2:
        for stop in stops:
            stop.arrival_sec = None


def _repair_isolated_time_outliers(stops: list[_ParsedStop]) -> None:
    events = [
        stop.departure_sec if stop.departure_sec is not None else stop.arrival_sec for stop in stops
    ]
    for index in range(1, len(events) - 1):
        previous = events[index - 1]
        current = events[index]
        following = events[index + 1]
        if previous is None or current is None or following is None:
            continue
        # Only repair a single value that lies outside two ordered neighbors.
        # This catches provider spreadsheet typos without smoothing valid long
        # express segments or changing an entire inconsistent column.
        if following >= previous and not previous <= current <= following:
            repaired = round((previous + following) / 2)
            stop = stops[index]
            if stop.arrival_sec is not None:
                stop.arrival_sec = repaired
            if stop.departure_sec is not None:
                stop.departure_sec = repaired
            events[index] = repaired


def _normalize_trip_times(stops: list[_ParsedStop], fallback_segment_seconds: int) -> None:
    offset = 0
    previous: int | None = None
    for stop in stops:
        if stop.arrival_sec is not None:
            if (
                previous is not None
                and stop.arrival_sec + offset < previous - ROLLOVER_THRESHOLD_SECONDS
            ):
                offset += SECONDS_PER_DAY
            stop.arrival_sec += offset
            previous = stop.arrival_sec
        if stop.departure_sec is not None:
            if (
                previous is not None
                and stop.departure_sec + offset < previous - ROLLOVER_THRESHOLD_SECONDS
            ):
                offset += SECONDS_PER_DAY
            stop.departure_sec += offset
            previous = stop.departure_sec

    events = [
        stop.departure_sec if stop.departure_sec is not None else stop.arrival_sec for stop in stops
    ]
    known_deltas = [
        later - earlier
        for earlier, later in pairwise(events)
        if earlier is not None and later is not None and 0 < later - earlier <= 15 * 60
    ]
    typical_delta = (
        max(1, round(statistics.median(known_deltas)))
        if known_deltas
        else fallback_segment_seconds
    )

    for index, event in enumerate(events):
        if event is not None:
            continue
        previous_index = next(
            (candidate for candidate in range(index - 1, -1, -1) if events[candidate] is not None),
            None,
        )
        next_index = next(
            (
                candidate
                for candidate in range(index + 1, len(events))
                if events[candidate] is not None
            ),
            None,
        )
        if previous_index is not None and next_index is not None:
            previous_value = events[previous_index]
            next_value = events[next_index]
            assert previous_value is not None
            assert next_value is not None
            span = next_index - previous_index
            step = (next_value - previous_value) / span
            events[index] = round(previous_value + step * (index - previous_index))
        elif previous_index is not None:
            previous_value = events[previous_index]
            assert previous_value is not None
            events[index] = previous_value + typical_delta * (index - previous_index)
        elif next_index is not None:
            next_value = events[next_index]
            assert next_value is not None
            events[index] = next_value - typical_delta * (next_index - index)

    for index, stop in enumerate(stops):
        event = events[index]
        if stop.arrival_sec is None and stop.departure_sec is None:
            if index == len(stops) - 1:
                stop.arrival_sec = event
            else:
                stop.arrival_sec = event
                stop.departure_sec = event
        elif stop.arrival_sec is None:
            stop.arrival_sec = stop.departure_sec
    _repair_isolated_time_outliers(stops)


def _workbook_rows(path: Path) -> Iterator[tuple[int, dict[str, str]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise SourceFormatError(f"Unable to open KRIC XLSX: {error}") from error
    found_sheet = False
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            header: list[str] | None = None
            header_row_number = 0
            for row_number, row in enumerate(rows, start=1):
                candidate = [
                    HEADER_LOOKUP.get(_normalize_header(value), "") for value in row
                ]
                if REQUIRED_HEADERS.issubset(set(candidate)):
                    header = candidate
                    header_row_number = row_number
                    found_sheet = True
                    break
                if row_number >= 20:
                    break
            if header is None:
                continue
            for row_number, row in enumerate(rows, start=header_row_number + 1):
                values = [_cell_text(value) for value in row]
                record = {
                    name: values[index] if index < len(values) else ""
                    for index, name in enumerate(header)
                    if name
                }
                if any(record.values()):
                    yield row_number, record
    finally:
        workbook.close()
    if not found_sheet:
        raise SourceFormatError("KRIC workbook does not contain the standard timetable headers")


def workbook_effective_from(path: Path, fallback: date) -> date:
    dates: set[date] = set()
    for _, row in _workbook_rows(path):
        value = row.get("data_as_of", "")[:10]
        if not value:
            continue
        try:
            dates.add(date.fromisoformat(value))
        except ValueError:
            continue
    return max(dates, default=fallback)


def iter_kric_workbook(path: Path, dataset: KricDataset) -> Iterator[NormalizedStopTime]:
    for row_number, row in _workbook_rows(path):
        station_text = row.get("stations", "")
        departure_text = row.get("departures", "")
        if not station_text or not departure_text:
            continue
        station_items = parse_station_list(station_text)
        arrivals = parse_timed_stops(row.get("arrivals", ""))
        departures = parse_timed_stops(departure_text)
        stops = [
            _ParsedStop(
                source_code=source_code,
                station_name=station_name,
                arrival_sec=arrivals.get(source_code),
                departure_sec=departures.get(source_code),
            )
            for source_code, station_name in station_items
        ]
        _prefer_consistent_time_column(stops)
        _normalize_trip_times(stops, dataset.fallback_segment_seconds)
        direction = _direction(stops, dataset.station_order)
        train_number = row.get("train_number", "") or str(row_number)
        operation_type = row.get("operation_type", "").upper()
        express_type = (
            "EXPRESS" if any(label in operation_type for label in ("급행", "직통", "EXPRESS")) else "LOCAL"
        )
        origin = normalize_station_name(row.get("origin", "")) or stops[0].station_name
        destination = (
            normalize_station_name(row.get("destination", ""))
            or stops[-1].station_name
        )
        for service_code in service_codes(row.get("service", "")):
            trip_code = (
                f"{dataset.line_code}:{service_code}:{direction}:{train_number}:{row_number}"
            )
            for stop_sequence, stop in enumerate(stops, start=1):
                yield NormalizedStopTime(
                    line_code=dataset.line_code,
                    station_code=_stable_station_code(dataset.line_code, stop.station_name),
                    station_name=stop.station_name,
                    service_code=service_code,
                    direction=direction,
                    trip_code=trip_code,
                    train_number=train_number,
                    stop_sequence=stop_sequence,
                    arrival_sec=stop.arrival_sec,
                    departure_sec=stop.departure_sec,
                    origin_name=origin,
                    destination_name=destination,
                    express_type=express_type,
                    source_row_number=row_number,
                )


class KricStandardTimetableCollector(TransitSourceCollector):
    def __init__(self, settings: Settings, dataset: KricDataset) -> None:
        self._settings = settings
        self._dataset = dataset
        self._downloader = ArtifactDownloader(
            timeout_seconds=settings.http_timeout_seconds,
            max_retries=settings.http_max_retries,
            timezone=settings.timezone,
        )

    async def inspect(self) -> SourceMetadata:
        return SourceMetadata(
            code=self._dataset.source_code,
            name=self._dataset.source_name,
            kind=SourceKind.DOWNLOAD_FILE,
            coverage=CoverageStatus.FULL_STATIC,
            access_status=(
                SourceAccessStatus.ENABLED
                if self._dataset.download_url
                else SourceAccessStatus.PENDING_REVIEW
            ),
            base_url=KRIC_DATASET_PAGE,
            stores_raw_artifact=True,
            uses_private_api=False,
            notes=(
                f"KRIC 표준데이터 운행정보 공개 XLSX (dataset id={self._dataset.dataset_id}); "
                "API 키 불필요"
            ),
        )

    async def collect(self, previous: CollectedArtifact | None = None) -> CollectionResult | None:
        today = datetime.now(tz=self._settings.timezone).date()
        artifact = await self._downloader.download(
            source_code=self._dataset.source_code,
            url=self._dataset.download_url,
            recorded_url=KRIC_DATASET_PAGE,
            destination=(
                self._settings.raw_data_dir
                / self._dataset.source_code
                / f"kric-{self._dataset.dataset_id}-{today.isoformat()}.xlsx"
            ),
            conditional=ConditionalRequest(
                etag=previous.etag if previous else None,
                last_modified=previous.last_modified if previous else None,
            ),
        )
        if artifact is None or (previous and previous.checksum == artifact.checksum):
            return None
        return CollectionResult(
            artifact=artifact,
            effective_from=workbook_effective_from(
                artifact.path, fallback=self._dataset.effective_from
            ),
        )

    def iter_records(self, artifact: CollectedArtifact) -> Iterator[NormalizedStopTime]:
        yield from iter_kric_workbook(artifact.path, self._dataset)

    async def validate(self, artifact: CollectedArtifact) -> ValidationResult:
        validation = validate_normalized_stop_times(
            self.iter_records(artifact), required_line_codes={self._dataset.line_code}
        )
        row_count = 0
        for _ in self.iter_records(artifact):
            row_count += 1
        if row_count >= self._dataset.minimum_record_count:
            return validation
        from metro_collector.domain.enums import QualitySeverity
        from metro_collector.domain.models import QualityResult

        minimum_check = QualityResult(
            rule_code="EXPECTED_MINIMUM_RECORD_COUNT",
            severity=QualitySeverity.FATAL,
            passed=False,
            affected_count=self._dataset.minimum_record_count - row_count,
            details={
                "parsed_rows": row_count,
                "minimum_rows": self._dataset.minimum_record_count,
            },
        )
        return ValidationResult(passed=False, checks=(*validation.checks, minimum_check))

    async def activate(self, version_id: str) -> ActivationResult:
        return ActivationResult(
            activated=False,
            version_id=version_id,
            reason="Activation is transactionally performed by TimetableIngestionService",
        )


def build_kric_standard_collectors(settings: Settings) -> list[KricStandardTimetableCollector]:
    datasets = (
        KricDataset(
            source_code="SHINBUNDANG_WEB",
            source_name="신분당선 표준데이터 운행정보",
            dataset_id=15,
            line_code="SHINBUNDANG",
            effective_from=date(2026, 6, 18),
            download_url=settings.kric_shinbundang_timetable_download_url,
            station_order=(
                "신사",
                "논현",
                "신논현",
                "강남",
                "양재",
                "양재시민의숲",
                "청계산입구",
                "판교",
                "정자",
                "미금",
                "동천",
                "수지구청",
                "성복",
                "상현",
                "광교중앙",
                "광교",
            ),
        ),
        KricDataset(
            source_code="GIMPO_GOLD_WEB",
            source_name="김포골드라인 표준데이터 운행정보",
            dataset_id=903,
            line_code="GIMPO_GOLD",
            effective_from=date(2026, 6, 16),
            download_url=settings.kric_gimpo_gold_timetable_download_url,
            station_order=(
                "양촌",
                "구래",
                "마산",
                "장기",
                "운양",
                "걸포북변",
                "사우(김포시청)",
                "풍무",
                "고촌",
                "김포공항",
            ),
        ),
    )
    return [KricStandardTimetableCollector(settings, dataset) for dataset in datasets]
