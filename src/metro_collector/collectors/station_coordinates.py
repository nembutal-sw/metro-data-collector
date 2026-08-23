from __future__ import annotations

from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from metro_collector.collectors.base import TransitSourceCollector
from metro_collector.collectors.http import ArtifactDownloader, ConditionalRequest
from metro_collector.config import Settings
from metro_collector.domain.enums import (
    CoverageStatus,
    QualitySeverity,
    SourceAccessStatus,
    SourceKind,
)
from metro_collector.domain.models import (
    ActivationResult,
    CollectedArtifact,
    CollectionResult,
    NormalizedStationCoordinate,
    NormalizedStopTime,
    QualityResult,
    SourceMetadata,
    ValidationResult,
)
from metro_collector.exceptions import CollectorConfigurationError, SourceFormatError
from metro_collector.routing.travel_time import normalize_station_name

SOURCE_CODE = "KRIC_STATION_COORDINATES"
DATASET_PAGE = "https://data.kric.go.kr/rips/M_01_01/detail.do?id=32"

# A KRIC infrastructure line can be used by more than one passenger-service line.
# Priority keeps the operator's direct line row ahead of a shared-track fallback.
SOURCE_LINE_TARGETS: dict[str, tuple[tuple[str, int], ...]] = {
    "1호선": (("SEOUL_1", 100),),
    "경부선": (("SEOUL_1", 90), ("SUIN_BUNDANG", 60)),
    "경원선": (
        ("SEOUL_1", 90),
        ("GYEONGUI_JUNGANG", 80),
        ("GYEONGCHUN", 70),
        ("SUIN_BUNDANG", 60),
    ),
    "경인선": (("SEOUL_1", 90),),
    "장항선": (("SEOUL_1", 90),),
    "2호선": (("SEOUL_2", 100),),
    "3호선": (("SEOUL_3", 100),),
    "일산선": (("SEOUL_3", 90),),
    "4호선": (("SEOUL_4", 100),),
    "안산과천선": (("SEOUL_4", 90),),
    "진접선": (("SEOUL_4", 90),),
    "5호선": (("SEOUL_5", 100),),
    "6호선": (("SEOUL_6", 100),),
    "7호선": (("SEOUL_7", 100),),
    "도시철도 7호선": (("SEOUL_7", 100),),
    "8호선": (("SEOUL_8", 100),),
    "수도권 광역철도 8호선": (("SEOUL_8", 100),),
    "서울 도시철도 9호선": (("SEOUL_9", 100),),
    "수도권  도시철도 9호선": (("SEOUL_9", 100),),
    "인천지하철 1호선": (("INCHEON_1", 100),),
    "인천지하철 2호선": (("INCHEON_2", 100),),
    "인천국제공항선": (("AREX", 100),),
    "신분당선": (("SHINBUNDANG", 100),),
    "우이신설선": (("UI_SINSEOL", 100),),
    "수도권 경량도시철도 신림선": (("SILLIM", 100),),
    "김포도시철도": (("GIMPO_GOLD", 100),),
    "의정부": (("UIJEONGBU_LRT", 100),),
    "에버라인": (("EVERLINE", 100),),
    "경춘선": (("GYEONGCHUN", 100),),
    "경강선": (("GYEONGGANG", 100),),
    "경의중앙선": (("GYEONGUI_JUNGANG", 100),),
    "서해선": (("SEOHAE", 100),),
    "분당선": (("SUIN_BUNDANG", 100),),
    "수인선": (("SUIN_BUNDANG", 100),),
}

HEADER_ALIASES = {
    "station_code": {"역번호"},
    "station_name": {"역사명", "역명"},
    "line_name": {"노선명", "선명"},
    "station_name_en": {"영문역사명", "영문역명"},
    "latitude": {"역위도", "위도"},
    "longitude": {"역경도", "경도"},
    "data_basis_date": {"데이터기준일자"},
}
HEADER_LOOKUP = {
    "".join(alias.split()): canonical
    for canonical, aliases in HEADER_ALIASES.items()
    for alias in aliases
}
REQUIRED_HEADERS = {"station_code", "station_name", "line_name", "latitude", "longitude"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value: Any) -> str:
    return "".join(_text(value).replace("\ufeff", "").split())


def _basis_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value).replace(".", "-").replace("/", "-")[:10]
    if not text:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        return None


def _workbook_rows(path: Path) -> Iterator[tuple[int, dict[str, Any]]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise SourceFormatError(f"Unable to open KRIC station workbook: {error}") from error
    found_header = False
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            columns: list[str] | None = None
            header_row = 0
            for row_number, row in enumerate(rows, start=1):
                candidate = [HEADER_LOOKUP.get(_header(value), "") for value in row]
                if REQUIRED_HEADERS.issubset(set(candidate)):
                    columns = candidate
                    header_row = row_number
                    found_header = True
                    break
                if row_number >= 20:
                    break
            if columns is None:
                continue
            for row_number, row in enumerate(rows, start=header_row + 1):
                record = {
                    name: row[index] if index < len(row) else None
                    for index, name in enumerate(columns)
                    if name
                }
                if any(value not in (None, "") for value in record.values()):
                    yield row_number, record
    finally:
        workbook.close()
    if not found_header:
        raise SourceFormatError("KRIC station workbook does not contain coordinate headers")


def scan_station_coordinates(
    path: Path,
) -> tuple[list[NormalizedStationCoordinate], list[str]]:
    records: list[NormalizedStationCoordinate] = []
    errors: list[str] = []
    for row_number, row in _workbook_rows(path):
        source_line = _text(row.get("line_name"))
        targets = SOURCE_LINE_TARGETS.get(source_line)
        if not targets:
            continue
        station_name = normalize_station_name(_text(row.get("station_name")))
        station_code = _text(row.get("station_code"))
        if not station_name or not station_code:
            errors.append(f"Missing station identity at row {row_number}")
            continue
        raw_latitude = row.get("latitude")
        raw_longitude = row.get("longitude")
        if raw_latitude is None or raw_longitude is None:
            errors.append(f"Missing coordinate at row {row_number}")
            continue
        try:
            latitude = float(raw_latitude)
            longitude = float(raw_longitude)
        except (TypeError, ValueError):
            errors.append(f"Invalid coordinate at row {row_number}")
            continue
        # The source workbook is nationwide, but every mapped target above is in the
        # Seoul metropolitan operating area. A broad box catches swapped or mistyped
        # values while leaving room for the outer ends of Line 1 and Gyeonggang Line.
        if not (36.5 <= latitude <= 38.5 and 126.0 <= longitude <= 128.5):
            errors.append(
                f"Coordinate outside metropolitan bounds at row {row_number}: "
                f"{latitude}, {longitude}"
            )
            continue
        english_name = _text(row.get("station_name_en")) or None
        for line_code, priority in targets:
            records.append(
                NormalizedStationCoordinate(
                    line_code=line_code,
                    source_line_name=source_line,
                    source_station_code=station_code,
                    station_name=station_name,
                    station_name_en=english_name,
                    latitude=latitude,
                    longitude=longitude,
                    data_basis_date=_basis_date(row.get("data_basis_date")),
                    priority=priority,
                    source_row_number=row_number,
                )
            )
    return records, errors


def iter_station_coordinates(path: Path) -> Iterator[NormalizedStationCoordinate]:
    records, _ = scan_station_coordinates(path)
    yield from records


class KricStationCoordinateCollector(TransitSourceCollector):
    def __init__(self, settings: Settings) -> None:
        self._settings = settings
        self._downloader = ArtifactDownloader(
            timeout_seconds=settings.http_timeout_seconds,
            max_retries=settings.http_max_retries,
            timezone=settings.timezone,
        )

    async def inspect(self) -> SourceMetadata:
        return SourceMetadata(
            code=SOURCE_CODE,
            name="국가철도공단 도시·광역철도 역사 위치",
            kind=SourceKind.DOWNLOAD_FILE,
            coverage=CoverageStatus.PARTIAL,
            access_status=(
                SourceAccessStatus.ENABLED
                if self._settings.kric_station_coordinates_download_url
                else SourceAccessStatus.PENDING_REVIEW
            ),
            base_url=DATASET_PAGE,
            stores_raw_artifact=True,
            uses_private_api=False,
            notes="공식 공개 XLSX dataset id=32; API 키와 REST 호출 한도 불필요",
        )

    async def collect(self, previous: CollectedArtifact | None = None) -> CollectionResult | None:
        url = self._settings.kric_station_coordinates_download_url
        if not url:
            raise CollectorConfigurationError(
                "KRIC_STATION_COORDINATES_DOWNLOAD_URL is not configured"
            )
        today = datetime.now(tz=self._settings.timezone).date()
        artifact = await self._downloader.download(
            source_code=SOURCE_CODE,
            url=url,
            recorded_url=DATASET_PAGE,
            destination=(
                self._settings.raw_data_dir
                / SOURCE_CODE
                / f"kric-station-coordinates-{today.isoformat()}.xlsx"
            ),
            conditional=ConditionalRequest(
                etag=previous.etag if previous else None,
                last_modified=previous.last_modified if previous else None,
            ),
        )
        # Re-apply an unchanged workbook because timetable ingestion can add new
        # station rows after the previous coordinate synchronization.
        if artifact is None or (previous and previous.checksum == artifact.checksum):
            if previous is None:
                return None
            artifact = previous
        basis_dates = [
            record.data_basis_date
            for record in iter_station_coordinates(artifact.path)
            if record.data_basis_date is not None
        ]
        return CollectionResult(
            artifact=artifact,
            effective_from=max(basis_dates, default=date(2024, 12, 31)),
        )

    def iter_coordinates(
        self, artifact: CollectedArtifact
    ) -> Iterator[NormalizedStationCoordinate]:
        yield from iter_station_coordinates(artifact.path)

    def iter_records(self, artifact: CollectedArtifact) -> Iterator[NormalizedStopTime]:
        del artifact
        raise TypeError("Station-coordinate artifacts do not contain timetable records")
        yield

    async def validate(self, artifact: CollectedArtifact) -> ValidationResult:
        records: list[NormalizedStationCoordinate] = []
        row_errors: list[str] = []
        format_errors: list[str] = []
        try:
            records, row_errors = scan_station_coordinates(artifact.path)
        except SourceFormatError as error:
            format_errors.append(str(error))
        count = len(records)
        row_error_limit = max(5, round((count + len(row_errors)) * 0.02))
        checks = (
            QualityResult(
                rule_code="SOURCE_FORMAT",
                severity=QualitySeverity.FATAL,
                passed=not format_errors,
                affected_count=len(format_errors),
                details={"sample_errors": format_errors[:20]},
            ),
            QualityResult(
                rule_code="VALID_COORDINATE_RATIO",
                severity=QualitySeverity.FATAL,
                passed=len(row_errors) <= row_error_limit,
                affected_count=len(row_errors),
                details={
                    "allowed_invalid_rows": row_error_limit,
                    "sample_errors": row_errors[:20],
                },
            ),
            QualityResult(
                rule_code="EXPECTED_COORDINATE_COUNT",
                severity=QualitySeverity.FATAL,
                passed=count >= 650,
                affected_count=0 if count >= 650 else 650 - count,
                details={"parsed_target_records": count},
            ),
        )
        return ValidationResult(passed=all(check.passed for check in checks), checks=checks)

    async def activate(self, version_id: str) -> ActivationResult:
        return ActivationResult(
            activated=False,
            version_id=version_id,
            reason="Station coordinates are applied by StationCoordinateIngestionService",
        )
