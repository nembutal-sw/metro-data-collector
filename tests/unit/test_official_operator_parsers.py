from datetime import date

from lxml import html  # type: ignore[import-untyped]
from openpyxl import Workbook  # type: ignore[import-untyped]

from metro_collector.collectors.official_operators import (
    StationHtmlDataset,
    _hourly_times,
    _synthesize_trips,
    iter_arex_workbook,
)


def test_hourly_html_table_parses_service_hours_past_midnight() -> None:
    table = html.fromstring(
        """
        <table>
          <tr><th>시</th><th>분</th></tr>
          <tr><td>5</td><td>01 12 59</td></tr>
          <tr><td>24</td><td>03 15</td></tr>
        </table>
        """
    )
    assert _hourly_times(table, hour_column=0, minute_column=1) == [
        5 * 3600 + 60,
        5 * 3600 + 12 * 60,
        5 * 3600 + 59 * 60,
        24 * 3600 + 3 * 60,
        24 * 3600 + 15 * 60,
    ]


def test_station_event_synthesis_keeps_middle_station_short_turn() -> None:
    dataset = StationHtmlDataset(
        source_code="TEST",
        source_name="test",
        line_code="TEST",
        effective_from=date(2026, 1, 1),
        station_order=("A", "B", "C"),
        page_urls=(),
        parser_kind="TEST",
        forward_segment_seconds=(60, 60),
        reverse_segment_seconds=(60, 60),
    )
    events = {
        ("WEEKDAY", "DOWN", 0): [5 * 3600],
        ("WEEKDAY", "DOWN", 1): [5 * 3600 + 60, 6 * 3600],
        ("WEEKDAY", "DOWN", 2): [],
    }
    trips = _synthesize_trips(events, dataset, "WEEKDAY", "DOWN")
    assert [trip.times for trip in trips] == [
        [(0, 5 * 3600), (1, 5 * 3600 + 60), (2, 5 * 3600 + 120)],
        [(1, 6 * 3600), (2, 6 * 3600 + 60)],
    ]


def test_arex_official_matrix_parser_skips_pass_through_cells(tmp_path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "평일"
    worksheet.cell(4, 2, "A1001")
    worksheet.cell(13, 1, "서울")
    worksheet.cell(14, 2, "05:00:00")
    worksheet.cell(15, 1, "공덕")
    worksheet.cell(15, 2, "---")
    worksheet.cell(16, 2, "05:04:00")
    worksheet.cell(17, 1, "홍대입구")
    worksheet.cell(17, 2, "05:07:00")
    worksheet.cell(18, 2, "05:07:30")
    path = tmp_path / "arex.xlsx"
    workbook.save(path)

    records = list(iter_arex_workbook(path))
    assert [record.station_name for record in records] == ["서울", "홍대입구"]
    assert [record.stop_sequence for record in records] == [1, 2]
    assert records[0].express_type == "EXPRESS"
    assert records[1].arrival_sec == 5 * 3600 + 7 * 60
